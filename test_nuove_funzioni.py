"""Test delle funzioni aggiunte: filigrana/protezione PDF, modello Excel batch,
manifest aggiornamenti, anteprima a copia singola."""

from __future__ import annotations

import json
import sys
import tempfile
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parent))

import app  # noqa: E402

ROOT = Path(__file__).resolve().parent


def _templates() -> list[app.TemplateFile]:
    templates, _ = app.discover_templates(ROOT / "templates")
    return [t for t in templates if t.path.suffix.lower() in {".docx", ".xlsx"}][:2]


def test_metadata_watermark_and_protection(tmp: Path) -> None:
    selected = _templates()
    plain = tmp / "plain.pdf"
    app.build_pdf(plain, "Mario Rossi", "01/10/2026", "SICUREZZA", "", "", selected)
    meta = PdfReader(str(plain)).metadata
    assert meta.title == "Dossier formazione - Mario Rossi", meta.title
    assert "SICUREZZA" in (meta.subject or "")

    marked = tmp / "marked.pdf"
    app.build_pdf(marked, "Mario Rossi", "01/10/2026", "SICUREZZA", "", "", selected,
                  watermark="COPIA CONTROLLATA", protect=True)
    reader = PdfReader(str(marked))
    assert reader.is_encrypted
    assert reader.decrypt("") != 0, "il PDF deve aprirsi senza password"
    assert len(reader.pages) == len(PdfReader(str(plain)).pages)
    assert "COPIA CONTROLLATA" in (reader.pages[-1].extract_text() or "")
    perms = reader.user_access_permissions
    assert perms & app_perm("PRINT") and not perms & app_perm("MODIFY")
    assert not list(tmp.glob("*.finalizing.pdf"))


def app_perm(name: str) -> int:
    from pypdf.constants import UserAccessPermissions
    return int(getattr(UserAccessPermissions, name))


def test_single_copy_preview(tmp: Path) -> None:
    selected = [replace(t, copies=3) for t in _templates()]
    single = [replace(t, copies=1) for t in selected]
    assert all(t.copies == 1 for t in single)
    out = tmp / f"{app.PREVIEW_PREFIX}test.pdf"
    total = app.build_pdf(out, "Mario Rossi", "01/10/2026", "X", "", "", single,
                          watermark="ANTEPRIMA")
    assert total == len(single)


def test_batch_template_roundtrip(tmp: Path) -> None:
    target = tmp / "modello.xlsx"
    app.write_batch_template(target, ["MAGAZZINO", "SICUREZZA"], italian=True)
    wb = load_workbook(target)
    assert wb.sheetnames[0] == "Dipendenti"
    assert wb["Reparti"].sheet_state == "hidden"
    wb.close()
    # Rilettura con lo stesso parser dell'import
    fake = app.FormazioniApp.__new__(app.FormazioniApp)
    fake.language = {}
    rows = app.FormazioniApp._parse_batch_file(fake, target)
    assert len(rows) == 1, rows
    assert rows[0]["Nome"] == "Mario Rossi"
    assert rows[0]["Reparto"] == "MAGAZZINO"
    assert len(rows[0]["Data"]) == 10 and rows[0]["Data"][2] == "/"

    english = tmp / "template_en.xlsx"
    app.write_batch_template(english, [], italian=False)
    rows = app.FormazioniApp._parse_batch_file(fake, english)
    assert set(rows[0]) >= {"Nome", "Data", "Reparto"}


def test_update_manifest(tmp: Path) -> None:
    assert app.parse_version("2.10.0") > app.parse_version("2.9.9")
    assert app.parse_version("v2.1") < app.parse_version("2.1.1")
    assert app.fetch_update_manifest("") is None
    (tmp / "version.json").write_text(json.dumps({"version": "9.0.0", "notes": "x"}),
                                      encoding="utf-8")
    manifest = app.fetch_update_manifest(str(tmp))
    assert manifest["version"] == "9.0.0"
    assert manifest["url"].endswith("FormazioniPZZ_Setup.exe")
    shipped = json.loads((ROOT / "release" / "version.json").read_text(encoding="utf-8"))
    assert shipped["version"] == app.APP_VERSION, "aggiorna release/version.json"

    # Risposta dell'API GitHub "releases/latest", senza rete
    import io
    import urllib.request
    github = {"tag_name": "v9.1.0", "name": "Formazioni PZZ 9.1.0",
              "html_url": "https://github.com/x/y/releases/tag/v9.1.0",
              "assets": [{"name": "FormazioniPZZ.exe", "browser_download_url": "https://e/portable"},
                         {"name": "FormazioniPZZ_Setup.exe", "browser_download_url": "https://e/setup"}]}
    original = urllib.request.urlopen
    urllib.request.urlopen = lambda *_a, **_k: io.BytesIO(json.dumps(github).encode())
    try:
        manifest = app.fetch_update_manifest(app.DEFAULT_UPDATE_SOURCE)
    finally:
        urllib.request.urlopen = original
    assert manifest == {"version": "9.1.0", "url": "https://e/setup",
                        "notes": "Formazioni PZZ 9.1.0"}, manifest


def test_placeholder_scan_and_hash_cache(tmp: Path) -> None:
    from docx import Document

    with_ph, without = tmp / "A_1_AAA.docx", tmp / "B_1_BBB.docx"
    doc = Document()
    paragraph = doc.add_paragraph("Ciao ")
    paragraph.add_run("*no")
    paragraph.add_run("me*")  # segnaposto spezzato in due run
    doc.save(with_ph)
    doc = Document()
    doc.add_paragraph("Regolamento senza dati personali")
    doc.save(without)
    assert app.template_has_placeholders(with_ph)
    assert not app.template_has_placeholders(without)
    assert app.template_has_placeholders(tmp / "x.doc"), ".doc: non verificabile"
    digest = app.compute_template_hash(without)
    assert app._HASH_CACHE[str(without)][2] == digest
    assert app.compute_template_hash(without) == digest


def _jobs(tmp: Path, count: int) -> list[app.DossierJob]:
    selected = _templates()
    return [app.DossierJob(tmp / f"d{i}.pdf", f"Persona {i}", "01/10/2026", "X", selected)
            for i in range(count)]


def test_batch_engine(tmp: Path) -> None:
    jobs = _jobs(tmp, 3)
    results = app.build_pdfs(jobs, watermark="BOZZA")
    assert results == [None, None, None], results
    for job in jobs:
        reader = PdfReader(str(job.output_path))
        assert reader.metadata.title == f"Dossier formazione - {job.employee_name}"
        assert "BOZZA" in (reader.pages[-1].extract_text() or "")
    if app._office_command():
        # I moduli senza segnaposto finiscono nella cache e vengono riusati
        static = [t for t in jobs[0].templates if not app.template_has_placeholders(t.path)]
        for t in static:
            assert app._cached_pdf_path(t.path).exists()


def test_parallel_engine_without_office(tmp: Path) -> None:
    original = app._office_command
    app._office_command = lambda: None  # forza il motore ReportLab (multi-processo)
    try:
        jobs = _jobs(tmp, 4)
        jobs.append(app.DossierJob(tmp / "missing.pdf", "Errore", "01/10/2026", "X",
                                   [app.TemplateFile(tmp / "NON_1_ESI.docx", "NON", 1, "ESI")]))
        results = app.build_pdfs(jobs)
        assert results[:4] == [None] * 4, results
        assert isinstance(results[4], Exception)
        assert all(j.output_path.exists() for j in jobs[:4])
    finally:
        app._office_command = original


def main() -> int:
    failures = 0
    for test in (test_metadata_watermark_and_protection, test_single_copy_preview,
                 test_batch_template_roundtrip, test_update_manifest,
                 test_placeholder_scan_and_hash_cache, test_batch_engine,
                 test_parallel_engine_without_office):
        with tempfile.TemporaryDirectory() as tmp:
            try:
                test(Path(tmp))
                print(f"OK    {test.__name__}")
            except Exception as exc:  # noqa: BLE001
                failures += 1
                print(f"FAIL  {test.__name__}: {exc!r}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
