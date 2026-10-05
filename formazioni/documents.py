"""Accesso a Word/Excel (import ritardati) e sostituzione dei segnaposto."""

from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path
from typing import Any


# python-docx, openpyxl e reportlab.platypus pesano quasi un secondo all'avvio:
# si importano alla prima generazione, non all'apertura della finestra.
def Document(*args, **kwargs):  # noqa: N802 - stesso nome dell'API di python-docx
    from docx import Document as _document

    return _document(*args, **kwargs)


def load_workbook(*args, **kwargs):
    from openpyxl import load_workbook as _load_workbook

    return _load_workbook(*args, **kwargs)


def replace_placeholders(value: object, employee_name: str, entry_date: str) -> object:
    if not isinstance(value, str):
        return value
    value = re.sub(r"\*nome\*", employee_name, value, flags=re.IGNORECASE)
    return re.sub(r"\*data\*", entry_date, value, flags=re.IGNORECASE)


# I segnaposto si sostituiscono solo nei nodi di testo (w:t, t, a:t) dell'XML.
# python-docx e openpyxl riscrivevano invece run e celle interi: il setter run.text
# cancella le immagini nella stessa run (i loghi delle intestazioni) e openpyxl,
# risalvando la cartella, perde immagini dell'intestazione, forme e caselle di testo.
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_WORD_PART = re.compile(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml")
_SHEET_PART = re.compile(r"xl/(sharedStrings|worksheets/sheet\d+|drawings/drawing\d+)\.xml")
_SHEET_HEADERS = ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter")


def _replace_in_text_nodes(nodes: list, employee_name: str, entry_date: str) -> bool:
    """Sostituisce i segnaposto nel testo di una sequenza di nodi (le run di un
    paragrafo). Un segnaposto spezzato su piu' nodi finisce nel primo; gli altri
    elementi (immagini, formattazione) restano dove sono."""
    original = [node.text or "" for node in nodes]
    texts = list(original)
    joined = "".join(texts)
    matches = list(re.finditer(r"\*(nome|data)\*", joined, flags=re.IGNORECASE))
    if not matches:
        return False
    starts = []
    offset = 0
    for text in original:
        starts.append(offset)
        offset += len(text)

    def locate(position: int, is_end: bool) -> tuple[int, int]:
        # Nodo che contiene il carattere in "position" (o, per la fine, quello prima)
        for index in range(len(original) - 1, -1, -1):
            if original[index] and (starts[index] < position if is_end else starts[index] <= position):
                return index, position - starts[index]
        return 0, position

    # Dall'ultimo al primo, cosi' le posizioni dei segnaposto precedenti restano valide
    for match in reversed(matches):
        value = employee_name if match.group(1).lower() == "nome" else entry_date
        first, first_offset = locate(match.start(), False)
        last, last_offset = locate(match.end(), True)
        if first == last:
            texts[first] = texts[first][:first_offset] + value + texts[first][last_offset:]
        else:
            texts[first] = texts[first][:first_offset] + value
            for index in range(first + 1, last):
                texts[index] = ""
            texts[last] = texts[last][last_offset:]
    for node, before, text in zip(nodes, original, texts):
        if before != text:
            node.text = text
            if node.tag == f"{{{_W}}}t":
                node.set(_XML_SPACE, "preserve")
    return True


def _grouped_text_nodes(root, container_tag: str, text_tag: str) -> list[list]:
    """Nodi di testo raggruppati per contenitore piu' vicino (paragrafo o stringa):
    il testo di una casella di testo dentro un paragrafo resta un gruppo a parte."""
    groups: dict[Any, list] = {}
    for container in root.iter(container_tag):
        groups[container] = []
    for node in root.iter(text_tag):
        parent = node.getparent()
        while parent is not None and parent.tag != container_tag:
            parent = parent.getparent()
        if parent is not None:
            groups[parent].append(node)
    return [nodes for nodes in groups.values() if nodes]


def _replace_in_xml_root(root, employee_name: str, entry_date: str) -> bool:
    changed = False
    for container, text in ((f"{{{_W}}}p", f"{{{_W}}}t"),
                            (f"{{{_S}}}si", f"{{{_S}}}t"),
                            (f"{{{_S}}}is", f"{{{_S}}}t"),
                            (f"{{{_A}}}p", f"{{{_A}}}t")):
        for nodes in _grouped_text_nodes(root, container, text):
            changed |= _replace_in_text_nodes(nodes, employee_name, entry_date)
    for name in _SHEET_HEADERS:
        for node in root.iter(f"{{{_S}}}{name}"):
            changed |= _replace_in_text_nodes([node], employee_name, entry_date)
    return changed


def replace_docx_placeholders(document, employee_name: str, entry_date: str) -> None:
    """Sostituisce i segnaposto in un Document di python-docx (corpo, tabelle,
    caselle di testo, tutte le intestazioni e i pie' di pagina)."""
    for part in document.part.package.iter_parts():
        if _WORD_PART.fullmatch(str(part.partname).lstrip("/")) and hasattr(part, "element"):
            _replace_in_xml_root(part.element, employee_name, entry_date)


def fill_office_placeholders(path: Path, employee_name: str, entry_date: str) -> bool:
    """Sostituisce i segnaposto in un .docx/.xlsx modificando solo l'XML del testo.
    Tutte le altre parti del file (immagini, intestazioni, forme, stili) vengono
    copiate byte per byte. Ritorna True se il file e' stato modificato."""
    from lxml import etree

    pattern = _WORD_PART if path.suffix.lower() == ".docx" else _SHEET_PART
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    changed: dict[str, bytes] = {}
    for info, data in entries:
        if not pattern.fullmatch(info.filename) or b"*" not in data:
            continue
        root = etree.fromstring(data, etree.XMLParser(resolve_entities=False, huge_tree=True))
        if _replace_in_xml_root(root, employee_name, entry_date):
            changed[info.filename] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if not changed:
        return False
    temp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(temp, "w") as output:
        for info, data in entries:
            output.writestr(info, changed.get(info.filename, data), compress_type=info.compress_type)
    os.replace(temp, path)
    return True

