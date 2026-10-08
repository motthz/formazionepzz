"""Modello Excel per l'import di piu' dipendenti."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from .config import ALL_DEPARTMENT_NAMES


def write_batch_template(path: Path, departments: list[str], italian: bool = True) -> None:
    """Crea un file Excel pronto da compilare per la generazione multipla."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    headers = ["Nome", "Data", "Reparto"] if italian else ["Name", "Date", "Department"]
    rows_with_rules = 500
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dipendenti" if italian else "Employees"
    sheet.append(headers)
    example_dept = next((d for d in departments if d.upper() not in ALL_DEPARTMENT_NAMES),
                        departments[0] if departments else "TUTTI")
    # L'esempio sta nei commenti delle intestazioni, non in una riga: una riga di
    # esempio dimenticata nel file diventava un dossier in piu' nel batch.
    examples = ("Mario Rossi", date.today().strftime("%d/%m/%Y"), example_dept)
    label = "Esempio" if italian else "Example"
    for cell, example in zip(sheet[1], examples):
        cell.comment = Comment(f"{label}: {example}", "Formazioni PZZ")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B2A3D")
        cell.alignment = Alignment(vertical="center")
    for row in range(2, rows_with_rules + 2):
        sheet.cell(row=row, column=2).number_format = "DD/MM/YYYY"
    for column, width in zip("ABC", (34, 16, 28)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    date_rule = DataValidation(type="date", operator="greaterThan", formula1="DATE(1990,1,1)",
                               allow_blank=True, errorStyle="warning")
    date_rule.error = "Data non valida (gg/mm/aaaa)" if italian else "Invalid date (dd/mm/yyyy)"
    sheet.add_data_validation(date_rule)
    date_rule.add(f"B2:B{rows_with_rules + 1}")
    if departments:
        # Elenco reparti su un foglio nascosto: niente limite di 255 caratteri.
        lists = workbook.create_sheet("Reparti" if italian else "Departments")
        for index, dept in enumerate(departments, start=1):
            lists.cell(row=index, column=1, value=dept)
        lists.sheet_state = "hidden"
        dept_rule = DataValidation(type="list", allow_blank=True, errorStyle="warning",
                                   formula1=f"='{lists.title}'!$A$1:$A${len(departments)}")
        dept_rule.error = ("Reparto non in elenco. Per più reparti usa A+B."
                           if italian else "Department not listed. For several use A+B.")
        sheet.add_data_validation(dept_rule)
        dept_rule.add(f"C2:C{rows_with_rules + 1}")
    workbook.save(path)
