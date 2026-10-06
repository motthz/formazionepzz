"""Sostituzione di *nome* e *data* nei modelli Word/Excel senza perdere immagini,
loghi delle intestazioni o altre parti del file."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Cm
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from formazioni.documents import fill_office_placeholders, replace_docx_placeholders
from formazioni.office import _prepare_office_template


def _logo(folder: Path) -> Path:
    path = folder / "logo.png"
    Image.new("RGB", (60, 30), "red").save(path)
    return path


def _entries(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_word_keeps_header_logo_and_images(tmp: Path) -> None:
    logo = _logo(tmp)
    document = Document()
    header = document.sections[0].header.paragraphs[0]
    # Logo e segnaposto nella stessa run: il vecchio run.text = ... cancellava il logo
    header.add_run().add_picture(str(logo), width=Cm(2))
    header.runs[0].add_text("Dipendente *nome*")
    body = document.add_paragraph()
    body.add_run("Logo ").add_picture(str(logo), width=Cm(1))
    body.add_run(" firmato da *no")
    body.add_run("me* il *data*")
    source = tmp / "TUTTI_1_ABC.docx"
    document.save(source)
    before = _entries(source)

    prepared = _prepare_office_template(source, tmp / "out", "Mario Rossi", "01/09/2026")

    after = _entries(prepared)
    assert after.keys() == before.keys()
    for name, data in before.items():
        if name not in {"word/document.xml", "word/header1.xml"}:
            assert after[name] == data, f"{name} non deve cambiare"
    assert after["word/document.xml"].count(b"<w:drawing") == 1
    assert after["word/header1.xml"].count(b"<w:drawing") == 1
    result = Document(str(prepared))
    assert result.paragraphs[0].text == "Logo  firmato da Mario Rossi il 01/09/2026"
    assert result.sections[0].header.paragraphs[0].text == "Dipendente Mario Rossi"


def test_word_in_memory_replacement_keeps_images(tmp: Path) -> None:
    document = Document()
    run = document.add_paragraph().add_run("*DATA* ")
    run.add_picture(str(_logo(tmp)), width=Cm(1))
    replace_docx_placeholders(document, "Anna", "02/02/2026")
    assert document.paragraphs[0].text == "02/02/2026 "
    assert document.paragraphs[0]._p.xml.count("<w:drawing") == 1


def test_excel_keeps_images_and_fills_header(tmp: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Dipendente: *nome*"
    sheet["A2"] = "Assunto il *data*"
    sheet["B1"] = 42
    sheet.oddHeader.center.text = "Scheda di *nome*"
    sheet.add_image(SheetImage(str(_logo(tmp))), "C1")
    source = tmp / "TUTTI_1_XLS.xlsx"
    workbook.save(source)
    before = _entries(source)

    assert fill_office_placeholders(source, "Mario Rossi", "01/09/2026")

    after = _entries(source)
    assert after.keys() == before.keys()
    changed = {name for name in before if after[name] != before[name]}
    assert changed <= {"xl/sharedStrings.xml", "xl/worksheets/sheet1.xml"}
    from openpyxl import load_workbook

    result = load_workbook(source).active
    assert result["A1"].value == "Dipendente: Mario Rossi"
    assert result["A2"].value == "Assunto il 01/09/2026"
    assert result["B1"].value == 42
    assert "Scheda di Mario Rossi" in result.oddHeader.center.text


def test_file_without_placeholders_is_untouched(tmp: Path) -> None:
    document = Document()
    document.add_paragraph("Nessun segnaposto qui * solo asterischi *")
    source = tmp / "fisso.docx"
    document.save(source)
    before = source.read_bytes()
    assert not fill_office_placeholders(source, "Mario Rossi", "01/09/2026")
    assert source.read_bytes() == before


def test_excel_sheet_name_removed_from_header_footer(tmp: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro interno"
    sheet["A1"] = "Contenuto fisso"
    sheet.oddFooter.center.text = "&A"
    sheet.oddFooter.right.text = "Pagina &P - R&&D"
    sheet.oddHeader.left.text = "Scheda di *nome* (&A)"
    source = tmp / "TUTTI_1_REG.xlsx"
    workbook.save(source)

    assert fill_office_placeholders(source)  # senza nome e data: solo il nome del foglio
    from openpyxl import load_workbook

    result = load_workbook(source).active
    assert not result.oddFooter.center.text
    assert result.oddFooter.right.text == "Pagina &P - R&&D"
    assert result.oddHeader.left.text == "Scheda di *nome* ()"
    assert result["A1"].value == "Contenuto fisso"


def _add_field(paragraph, code: str, shown: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def char(kind: str):
        run = paragraph.add_run()
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        run._r.append(element)

    char("begin")
    instr = OxmlElement("w:instrText")
    instr.text = f" {code} "
    paragraph.add_run()._r.append(instr)
    char("separate")
    paragraph.add_run(shown)
    char("end")


def test_word_filename_field_removed_from_footer(tmp: Path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("Corpo del modulo")
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("Rev. 3 - ")
    _add_field(footer, "FILENAME \\* MERGEFORMAT", "MAGAZZINO_2_MAG.docx")
    footer.add_run(" - pag. ")
    _add_field(footer, "PAGE", "1")
    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), " FILENAME \\p ")
    footer._p.append(simple)
    source = tmp / "MAGAZZINO_2_MAG.docx"
    document.save(source)

    assert fill_office_placeholders(source)

    result = Document(str(source))
    xml = result.sections[0].footer.paragraphs[0]._p.xml
    assert "FILENAME" not in xml
    assert "PAGE" in xml
    assert result.sections[0].footer.paragraphs[0].text == "Rev. 3 -  - pag. 1"
    assert result.paragraphs[0].text == "Corpo del modulo"
